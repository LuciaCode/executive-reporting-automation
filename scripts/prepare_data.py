import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

def prepare_sales_data(input_csv, output_xlsx):
    # 1. Load the sales data
    df = pd.read_csv(input_csv)
    
    # 2. Create a new Workbook and Sheet
    wb = Workbook()
    ws = wb.active
    ws.title = "SalesData"

    # 3. Add data to the sheet
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)

    # 4. Format as a Table (Hard Requirement for Copilot)
    # Define the range (e.g., A1:D25)
    tab_range = f"A1:{chr(64 + len(df.columns))}{len(df) + 1}"
    tab = Table(displayName="Ventas_2024", ref=tab_range)

    # Add a style
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws.add_table(tab)

    # 5. Add Metadata Sheet for LLM Context
    metadata_ws = wb.create_sheet(title="Metadata")
    metadata_data = [
        ["Field", "Description", "Data Type"],
        ["Date", "Transaction date (YYYY-MM-DD)", "Date"],
        ["Product Category", "High-level grouping of services/hardware", "String"],
        ["Units Sold", "Quantity of items sold in the transaction", "Integer"],
        ["Revenue", "Total currency value of the sale (USD)", "Decimal"],
        [],
        ["Table Name", "Ventas_2024", ""],
        ["Context", "This table contains all validated sales records for the fiscal year 2024.", ""]
    ]
    
    for row in metadata_data:
        metadata_ws.append(row)

    # Save the file
    wb.save(output_xlsx)
    print(f"Successfully created: {output_xlsx} with Table 'Ventas_2024' and Metadata sheet.")

if __name__ == "__main__":
if len(sys.argv) < 3:
        print("Uso: python scripts/prepare_data.py <ruta_entrada_csv> <nombre_salida_xlsx>")
    else:
        input_path = sys.argv[1]
        output_path = sys.argv[2]
        prepare_sales_data(input_path, output_path)